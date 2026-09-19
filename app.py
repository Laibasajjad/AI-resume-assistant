import io
import json
import os
import re
from typing import Any, Dict, List, Optional, Tuple

import streamlit as st
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document


APP_TITLE = "AI Resume Assistant"
DEFAULT_MODEL = "gemini-3.8-flash"
SUPPORTED_TYPES = ["pdf", "docx", "txt", "md"]


ANALYSIS_SCHEMA = {
    "type": "object",
    "properties": {
        "ats_score": {"type": "number"},
        "score_breakdown": {
            "type": "object",
            "properties": {
                "keyword_match": {"type": "number"},
                "experience_relevance": {"type": "number"},
                "skills_match": {"type": "number"},
                "education_match": {"type": "number"},
                "achievement_impact": {"type": "number"},
                "format_readability": {"type": "number"},
            },
            "required": [
                "keyword_match",
                "experience_relevance",
                "skills_match",
                "education_match",
                "achievement_impact",
                "format_readability",
            ],
        },
        "candidate_summary": {"type": "string"},
        "matched_keywords": {"type": "array", "items": {"type": "string"}},
        "missing_keywords": {"type": "array", "items": {"type": "string"}},
        "partially_matched_keywords": {
            "type": "array",
            "items": {"type": "string"},
        },
        "strengths": {"type": "array", "items": {"type": "string"}},
        "shortcomings": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "issue": {"type": "string"},
                    "evidence": {"type": "string"},
                    "impact": {"type": "string"},
                    "fix": {"type": "string"},
                },
                "required": ["issue", "evidence", "impact", "fix"],
            },
        },
        "section_feedback": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "section": {"type": "string"},
                    "status": {"type": "string"},
                    "feedback": {"type": "string"},
                },
                "required": ["section", "status", "feedback"],
            },
        },
        "improvements": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "priority": {"type": "string"},
                    "change": {"type": "string"},
                    "reason": {"type": "string"},
                    "example": {"type": "string"},
                },
                "required": ["priority", "change", "reason", "example"],
            },
        },
        "rewritten_bullets": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "original_or_target": {"type": "string"},
                    "improved_version": {"type": "string"},
                    "why_better": {"type": "string"},
                },
                "required": ["original_or_target", "improved_version", "why_better"],
            },
        },
        "ats_warnings": {"type": "array", "items": {"type": "string"}},
        "final_action_plan": {"type": "array", "items": {"type": "string"}},
    },
    "required": [
        "ats_score",
        "score_breakdown",
        "candidate_summary",
        "matched_keywords",
        "missing_keywords",
        "partially_matched_keywords",
        "strengths",
        "shortcomings",
        "section_feedback",
        "improvements",
        "rewritten_bullets",
        "ats_warnings",
        "final_action_plan",
    ],
}


def get_api_key() -> Optional[str]:
    """Read the Gemini key from Streamlit secrets or an environment variable."""
    try:
        secret_key = st.secrets.get("GEMINI_API_KEY")
        if secret_key:
            return str(secret_key)
    except Exception:
        pass
    return os.getenv("GEMINI_API_KEY")


def get_model() -> str:
    try:
        configured = st.secrets.get("GEMINI_MODEL")
        if configured:
            return str(configured)
    except Exception:
        pass
    return os.getenv("GEMINI_MODEL", DEFAULT_MODEL)


def extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n\n".join(pages).strip()


def extract_docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    chunks: List[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                chunks.append(" | ".join(cells))

    return "\n".join(chunks).strip()


def extract_text(uploaded_file) -> Tuple[str, Optional[bytes]]:
    data = uploaded_file.getvalue()
    extension = uploaded_file.name.lower().rsplit(".", 1)[-1]

    if extension == "pdf":
        return extract_pdf_text(data), data
    if extension == "docx":
        return extract_docx_text(data), None
    if extension in {"txt", "md"}:
        return data.decode("utf-8", errors="replace").strip(), None

    raise ValueError("Unsupported file type.")


def clean_json_text(text: str) -> str:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
        text = re.sub(r"\s*```$", "", text)
    return text.strip()


def validate_result(result: Dict[str, Any]) -> Dict[str, Any]:
    required = [
        "ats_score",
        "score_breakdown",
        "candidate_summary",
        "matched_keywords",
        "missing_keywords",
        "partially_matched_keywords",
        "strengths",
        "shortcomings",
        "section_feedback",
        "improvements",
        "rewritten_bullets",
        "ats_warnings",
        "final_action_plan",
    ]
    missing = [key for key in required if key not in result]
    if missing:
        raise ValueError(f"Gemini response is missing fields: {', '.join(missing)}")

    result["ats_score"] = max(0, min(100, float(result["ats_score"])))
    return result


def analyze_resume(
    client: genai.Client,
    resume_text: str,
    job_description: str,
    pdf_bytes: Optional[bytes] = None,
    model: str = DEFAULT_MODEL,
) -> Dict[str, Any]:
    """Ask Gemini for a structured ATS-style comparison."""
    if not resume_text.strip() and not pdf_bytes:
        raise ValueError("The resume could not be read. Try a text-based PDF or DOCX.")

    system_instruction = """
You are an expert resume reviewer and ATS-oriented career assistant.

Analyze ONLY the information present in the supplied resume and job description.
Do not invent experience, skills, certifications, education, metrics, employers, or dates.
A keyword should be considered matched only when the resume actually supports it.
Distinguish exact matches from reasonable equivalents and mark the latter as partial.
Never recommend adding a skill merely because it is in the job description if the candidate
has no evidence of having that skill. Instead, recommend learning it or demonstrating it
elsewhere.

ATS SCORE RULE:
Produce an ATS-style compatibility score from 0 to 100. This is an estimate, not a score
from a real ATS vendor. Use this conceptual weighting:
- keyword_match: 25
- experience_relevance: 20
- skills_match: 20
- education_match: 10
- achievement_impact: 15
- format_readability: 10
The six component scores should each be 0-100. The overall ats_score should be consistent
with these weights.

For shortcomings, be concrete. Give evidence from the resume, explain why it matters for
this job description, and give a specific fix. For rewritten bullets, preserve truth and
do not fabricate numbers.

Return valid JSON matching the requested schema and no markdown outside the JSON.
"""

    user_text = f"""
JOB DESCRIPTION:
{job_description}

RESUME TEXT EXTRACTED BY THE APP:
{resume_text[:50000]}

TASK:
Compare the resume against the job description and return the complete structured analysis.
Include:
1. ATS-style score and six-part breakdown.
2. Exact matched, missing, and partially matched keywords/skills.
3. Strengths relevant to this job.
4. A complete list of important shortcomings.
5. Section-by-section feedback.
6. Prioritized improvements.
7. Truth-preserving rewritten resume bullets where useful.
8. ATS formatting/readability warnings.
9. A practical final action plan.

Do not judge the person's worth or guarantee interview/job outcomes.
"""

    contents: List[Any] = [system_instruction, user_text]

    # For PDFs, give Gemini the original document too. This helps when the PDF contains
    # layout information or text extraction misses parts of the document.
    if pdf_bytes:
        uploaded = client.files.upload(
            file=io.BytesIO(pdf_bytes),
            config=types.UploadFileConfig(mime_type="application/pdf"),
        )
        contents = [system_instruction, user_text, uploaded]

    response = client.models.generate_content(
        model=model,
        contents=contents,
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
            response_schema=ANALYSIS_SCHEMA,
        ),
    )

    raw = response.text
    if not raw:
        raise ValueError("Gemini returned an empty response.")

    result = json.loads(clean_json_text(raw))
    return validate_result(result)


def local_only_validation(resume_text: str, job_description: str) -> Dict[str, Any]:
    """Small deterministic smoke-test helper; not the production ATS score."""
    resume_words = set(re.findall(r"[a-zA-Z][a-zA-Z0-9+#.-]{2,}", resume_text.lower()))
    jd_words = set(re.findall(r"[a-zA-Z][a-zA-Z0-9+#.-]{2,}", job_description.lower()))
    overlap = sorted(resume_words & jd_words)
    return {
        "resume_chars": len(resume_text),
        "job_description_chars": len(job_description),
        "overlap_sample": overlap[:20],
    }


def display_analysis(result: Dict[str, Any]) -> None:
    score = result["ats_score"]
    st.subheader("ATS Compatibility Estimate")
    st.metric("ATS-style score", f"{score:.0f}/100")
    st.progress(int(score))

    st.caption(
        "This is an AI-generated compatibility estimate, not a score produced by a "
        "specific ATS vendor."
    )

    st.subheader("Score Breakdown")
    breakdown = result["score_breakdown"]
    labels = {
        "keyword_match": "Keyword match",
        "experience_relevance": "Experience relevance",
        "skills_match": "Skills match",
        "education_match": "Education match",
        "achievement_impact": "Achievement impact",
        "format_readability": "Format/readability",
    }
    for key, label in labels.items():
        value = max(0, min(100, float(breakdown[key])))
        st.write(f"**{label}: {value:.0f}/100**")
        st.progress(int(value))

    st.subheader("Candidate vs Job Summary")
    st.write(result["candidate_summary"])

    col1, col2, col3 = st.columns(3)
    with col1:
        st.markdown("### Matched")
        for item in result["matched_keywords"]:
            st.write(f"✓ {item}")
    with col2:
        st.markdown("### Partial")
        for item in result["partially_matched_keywords"]:
            st.write(f"~ {item}")
    with col3:
        st.markdown("### Missing")
        for item in result["missing_keywords"]:
            st.write(f"✗ {item}")

    st.subheader("Strengths")
    for item in result["strengths"]:
        st.write(f"• {item}")

    st.subheader("Shortcomings")
    for item in result["shortcomings"]:
        with st.expander(item["issue"]):
            st.markdown(f"**Evidence:** {item['evidence']}")
            st.markdown(f"**Impact:** {item['impact']}")
            st.markdown(f"**Fix:** {item['fix']}")

    st.subheader("Section-by-Section Review")
    for item in result["section_feedback"]:
        with st.expander(f"{item['section']} — {item['status']}"):
            st.write(item["feedback"])

    st.subheader("Prioritized Improvements")
    for item in result["improvements"]:
        with st.expander(f"{item['priority']}: {item['change']}"):
            st.write(f"**Why:** {item['reason']}")
            st.write(f"**Example:** {item['example']}")

    st.subheader("Truth-Preserving Bullet Improvements")
    if not result["rewritten_bullets"]:
        st.info("No specific bullet rewrites were necessary.")
    else:
        for item in result["rewritten_bullets"]:
            with st.expander(item["original_or_target"]):
                st.markdown(f"**Improved version:** {item['improved_version']}")
                st.markdown(f"**Why:** {item['why_better']}")

    st.subheader("ATS Warnings")
    if result["ats_warnings"]:
        for item in result["ats_warnings"]:
            st.warning(item)
    else:
        st.success("No major ATS formatting warnings were identified.")

    st.subheader("Final Action Plan")
    for index, item in enumerate(result["final_action_plan"], 1):
        st.write(f"{index}. {item}")

    with st.expander("View analysis JSON"):
        st.json(result)


def main() -> None:
    st.set_page_config(
        page_title=APP_TITLE,
        page_icon="📄",
        layout="wide",
    )

    st.title("📄 AI Resume Assistant")
    st.write(
        "Upload a resume, paste a job description, and get an ATS-style compatibility "
        "analysis with keyword gaps, shortcomings, and actionable improvements."
    )

    with st.sidebar:
        st.header("Settings")
        model = st.text_input("Gemini model", value=get_model())
        st.caption("Keep your Gemini API key in Streamlit Secrets or an environment variable.")
        st.markdown(
            "Supported resume files: **PDF, DOCX, TXT, MD**."
        )

    uploaded_file = st.file_uploader(
        "Upload your resume",
        type=SUPPORTED_TYPES,
        help="Use a text-based PDF or DOCX for the most reliable extraction.",
    )

    job_description = st.text_area(
        "Paste the job description",
        height=300,
        placeholder="Paste the complete job description here...",
    )

    analyze_button = st.button("🔍 Analyze Resume", type="primary", use_container_width=True)

    if analyze_button:
        if not uploaded_file:
            st.error("Please upload a resume first.")
            return
        if not job_description.strip():
            st.error("Please paste the job description first.")
            return

        api_key = get_api_key()
        if not api_key:
            st.error(
                "Gemini API key not found. Add GEMINI_API_KEY to your local "
                ".streamlit/secrets.toml or your deployment secrets."
            )
            return

        try:
            with st.spinner("Reading your resume and asking Gemini to analyze the match..."):
                resume_text, pdf_bytes = extract_text(uploaded_file)

                if not resume_text.strip() and not pdf_bytes:
                    raise ValueError(
                        "No readable text was extracted. Please use a text-based PDF, DOCX, TXT, or MD file."
                    )

                client = genai.Client(api_key=api_key)
                result = analyze_resume(
                    client=client,
                    resume_text=resume_text,
                    job_description=job_description,
                    pdf_bytes=pdf_bytes,
                    model=model.strip() or DEFAULT_MODEL,
                )

            st.success("Analysis complete.")
            display_analysis(result)

        except json.JSONDecodeError:
            st.error(
                "Gemini returned an invalid JSON response. Try again; if the issue persists, "
                "try a shorter resume/job description."
            )
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")


if __name__ == "__main__":
    main()
